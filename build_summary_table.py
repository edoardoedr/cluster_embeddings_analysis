"""
Raccoglie i risultati prodotti da `calculate_clusters_embedding_labels.py`
(uno per label, in Results/clustering_metrics_results_labels_<LABEL>/) e li
riassume in un'unica tabella riepilogativa: una riga per foundation model
(FM), colonne raggruppate per metrica (kNN@10, Centroid separation, ARI) e,
dentro ogni gruppo, una colonna per label (CD, AF, SB, STach).

Ogni cella è la media della metrica sui dataset (chapman_ningbo, ptbxl,
cod15, georgia), presa da 'clustering_metrics_summary_by_model.csv'.

Output: Results/summary_table.xlsx (header a due righe con celle unite).
"""
import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
RESULTS_DIR = os.path.join(SCRIPT_DIR, 'Results')
OUTPUT_FILE = os.path.join(RESULTS_DIR, 'summary_table.xlsx')

LABELS = ['CD', 'AF', 'SB', 'STach']

# (nome colonna sorgente in clustering_metrics_summary_by_model.csv, nome gruppo in tabella)
METRICS = [
    ('knn_label_agreement_k10_mean', 'kNN@10'),
    ('centroid_separation_mean', 'Centroid sep'),
    ('gmm_ari_mean', 'ARI'),
]

# Ordine e nomi delle righe (FM) come richiesto, mappati alle cartelle model_configuration reali
MODEL_ROWS = [
    ('ECG-FM', 'ECG-FM'),
    ('ECGFounder', 'ECGFounder'),
    ('HuBERT-ECG small', 'HuBERT-ECG_small'),
    ('HuBERT-ECG base', 'HuBERT-ECG_base'),
    ('HuBERT-ECG large', 'HuBERT-ECG_large'),
    ('ECG-JEPA', 'ECG-JEPA'),
]


def load_label_summary(label):
    csv_path = os.path.join(
        RESULTS_DIR, f'clustering_metrics_results_labels_{label}', 'clustering_metrics_summary_by_model.csv'
    )
    df = pd.read_csv(csv_path)
    return df.set_index('model_configuration')


def build_table():
    label_data = {label: load_label_summary(label) for label in LABELS}

    columns = pd.MultiIndex.from_tuples(
        [(group, label) for _, group in METRICS for label in LABELS],
        names=['metric', 'label']
    )
    row_names = [display_name for display_name, _ in MODEL_ROWS]
    table = pd.DataFrame(index=row_names, columns=columns, dtype=float)

    for display_name, model_conf in MODEL_ROWS:
        for metric_col, group in METRICS:
            for label in LABELS:
                summary = label_data[label]
                if model_conf in summary.index and metric_col in summary.columns:
                    table.loc[display_name, (group, label)] = summary.loc[model_conf, metric_col]

    return table


def write_xlsx(table):
    n_groups = len(METRICS)
    n_labels = len(LABELS)
    first_data_col = 2  # colonna 1 = nomi riga (FM)
    first_data_row = 3  # righe 1-2 = header a due livelli

    wb = Workbook()
    ws = wb.active
    ws.title = 'summary'

    # Intestazione colonna FM (righe 1-2 unite)
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
    fm_header = ws.cell(row=1, column=1, value='FM')
    fm_header.alignment = Alignment(horizontal='center', vertical='center')
    fm_header.font = Font(bold=True)

    # Riga 1: nome gruppo metrica (celle unite su n_labels colonne), con freccia "più alto è meglio"
    for i, (_, group) in enumerate(METRICS):
        start_col = first_data_col + i * n_labels
        end_col = start_col + n_labels - 1
        ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=end_col)
        cell = ws.cell(row=1, column=start_col, value=f'{group} ↑')
        cell.alignment = Alignment(horizontal='center')
        cell.font = Font(bold=True)

    # Riga 2: nome label per ogni colonna
    col = first_data_col
    for _ in METRICS:
        for label in LABELS:
            cell = ws.cell(row=2, column=col, value=label)
            cell.alignment = Alignment(horizontal='center')
            cell.font = Font(bold=True)
            col += 1

    # Righe dati
    for r, row_name in enumerate(table.index):
        ws.cell(row=first_data_row + r, column=1, value=row_name)
        col = first_data_col
        for metric_col, group in METRICS:
            for label in LABELS:
                value = table.loc[row_name, (group, label)]
                cell = ws.cell(row=first_data_row + r, column=col,
                                value=None if pd.isna(value) else round(float(value), 4))
                cell.alignment = Alignment(horizontal='center')
                col += 1

    # Larghezza colonne
    ws.column_dimensions['A'].width = 18
    for c in range(first_data_col, first_data_col + n_groups * n_labels):
        ws.column_dimensions[get_column_letter(c)].width = 10

    wb.save(OUTPUT_FILE)
    print(f'Tabella salvata in: {OUTPUT_FILE}')


def main():
    table = build_table()
    print(table.round(4))
    write_xlsx(table)


if __name__ == '__main__':
    main()
